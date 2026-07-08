"""
v4.0.2 PPT 生成脚本
基于 v4.0.1 PPT 克隆，更新版本号和 v4.0.2 变更内容。
沿用 v3.22.24 风格（用户偏好）。
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC = Path(r"C:\Users\Administrator\Desktop\资产盘点拍照工具-使用说明书-v4.0.1.pptx")
DST = Path(r"C:\Users\Administrator\Desktop\资产盘点拍照工具-使用说明书-v4.0.2.pptx")

# 先复制
shutil.copy2(SRC, DST)

prs = Presentation(DST)

# 全局文本替换映射
REPLACEMENTS = {
    "v4.0.1": "v4.0.2",
    "V4.0.1": "V4.0.2",
    "4.0.1": "4.0.2",
    "versionName = \"4.0.1\"": "versionName = \"4.0.2\"",
    "versionCode = 2": "versionCode = 3",
}


def replace_in_run(run):
    """对单个 run 执行文本替换"""
    if not run.text:
        return
    new_text = run.text
    for old, new in REPLACEMENTS.items():
        if old in new_text:
            new_text = new_text.replace(old, new)
    if new_text != run.text:
        run.text = new_text


def walk_shapes(shapes):
    """递归遍历所有 shape（含 group 内的）"""
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from walk_shapes(shape.shapes)
        else:
            yield shape


# 第一遍：全局版本号替换
for slide in prs.slides:
    for shape in walk_shapes(slide.shapes):
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                replace_in_run(run)


# 新增一页：v4.0.2 变更说明（插入到末尾）
# 使用空白版式
blank_layout = prs.slide_layouts[6]
new_slide = prs.slides.add_slide(blank_layout)

# 标题
left = Inches(0.5)
top = Inches(0.4)
width = Inches(9)
height = Inches(0.8)
txBox = new_slide.shapes.add_textbox(left, top, width, height)
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run = p.add_run()
run.text = "v4.0.2 更新内容（2026-07-08）"
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = RGBColor(0x21, 0x21, 0x21)
run.font.name = "Microsoft YaHei"

# 内容
content_top = Inches(1.4)
content_left = Inches(0.5)
content_width = Inches(9)
content_height = Inches(5.5)
content_box = new_slide.shapes.add_textbox(content_left, content_top, content_width, content_height)
ctf = content_box.text_frame
ctf.word_wrap = True

changes = [
    ("1. 水印定位修复", True),
    ("   • 新增 GMS 可用性检查，失败时回退 LocationManager 原生 GPS（支持华为等无 GMS 设备）", False),
    ("   • 进入相机界面即触发位置预热，避免首次拍照冷启动超时", False),
    ("   • 定位超时 8s → 15s，移除 (0,0) 伪造坐标兜底", False),
    ("   • 定位失败时水印显示「定位失败」，不写入伪造经纬度", False),
    ("   • 定位失败时 Snackbar 提示用户检查权限/GPS", False),
    ("", False),
    ("2. 命名规则恢复", True),
    ("   • 新增 4 段下拉命名规则配置（拍摄日期/客户名/地址+时间/空值）", False),
    ("   • 自动追加类型+序号，如 20260708-成都投资集团-和平区XX街123号1430-远景-01.jpg", False),
    ("   • 设置页新增命名规则卡片，含实时预览", False),
    ("   • DataStore 持久化，全 NONE 时回退 IMG_<timestamp>.jpg（向后兼容）", False),
    ("", False),
    ("3. 客户条目排版重设计", True),
    ("   • 按钮改为左侧 48dp 竖版操作列（计数Badge/拍照/查看已拍/备注）", False),
    ("   • 右侧文本区 weight=1f，借款人名/地址/备注全量展示", False),
    ("   • 查看已拍按钮始终显示，避免误以为无反应", False),
    ("", False),
    ("4. 查看已拍按钮修复", True),
    ("   • 空状态文案改为「{客户名} 暂无照片（本次会话）」", False),
    ("   • 新增「在系统相册中查看」按钮，跳转系统相册", False),
    ("   • 缩略图加载失败显示占位图", False),
    ("", False),
    ("5. 全量计数显示", True),
    ("   • 始终展示所有 5 个分类（count=0 灰色显示）", False),
    ("   • 格式：总计 N  远景0 近景2 内部1 瑕疵0 其他0", False),
]

for idx, (text, is_bold) in enumerate(changes):
    if idx == 0:
        p = ctf.paragraphs[0]
    else:
        p = ctf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(13) if not is_bold else Pt(15)
    run.font.bold = is_bold
    run.font.color.rgb = RGBColor(0x21, 0x96, 0xF3) if is_bold else RGBColor(0x21, 0x21, 0x21)
    run.font.name = "Microsoft YaHei"

prs.save(DST)
print(f"PPT saved: {DST}")
print(f"Total slides: {len(prs.slides)}")
