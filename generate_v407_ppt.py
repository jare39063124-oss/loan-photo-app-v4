"""
v4.0.7 PPT 生成脚本
基于 v4.0.5-new.pptx，修复第11/12/13/15页文字重叠 + 版本号更新至 v4.0.7 + 新增 v4.0.6 与 v4.0.7 变更页
"""
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

SRC = Path(r"C:\Users\Administrator\Desktop\资产盘点拍照工具-使用说明书-v4.0.5-new.pptx")
DST = Path(r"C:\Users\Administrator\Desktop\资产盘点拍照工具-使用说明书-v4.0.7.pptx")

if not SRC.exists():
    SRC = Path(r"C:\Users\Administrator\Desktop\资产盘点拍照工具-使用说明书-v4.0.2.pptx")

shutil.copy2(SRC, DST)
prs = Presentation(DST)

print(f"幻灯片尺寸: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} 英寸")
print(f"原始页数: {len(prs.slides)}")

# === 第一部分：全局版本号替换 v4.0.5 → v4.0.7 ===
REPLACEMENTS = {
    "v4.0.5": "v4.0.7",
    "V4.0.5": "V4.0.7",
    "4.0.5": "4.0.7",
}

def replace_in_run(run):
    if not run.text:
        return
    new_text = run.text
    for old, new in REPLACEMENTS.items():
        if old in new_text:
            new_text = new_text.replace(old, new)
    if new_text != run.text:
        run.text = new_text

def walk_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from walk_shapes(shape.shapes)
        else:
            yield shape

for slide in prs.slides:
    for shape in walk_shapes(slide.shapes):
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                replace_in_run(run)

# === 第二部分：修复第11/12/13/15页文字重叠 ===
def fix_overlap(slide, page_num):
    """修复指定页的文本框重叠：增大摘要文本框高度 + 调整下方文本框 top"""
    shapes_with_tf = []
    for shape in walk_shapes(slide.shapes):
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                left = shape.left / 914400 if shape.left else 0
                top = shape.top / 914400 if shape.top else 0
                w = shape.width / 914400 if shape.width else 0
                h = shape.height / 914400 if shape.height else 0
                line_count = 0
                max_size = 0
                for p in shape.text_frame.paragraphs:
                    p_text = p.text
                    if p_text.strip():
                        line_count += max(1, p_text.count('\n') + 1)
                    for r in p.runs:
                        if r.font.size:
                            max_size = max(max_size, r.font.size.pt)
                shapes_with_tf.append({
                    'shape': shape,
                    'left': left, 'top': top, 'w': w, 'h': h,
                    'line_count': line_count,
                    'max_size': max_size if max_size > 0 else 14,
                    'text': text[:50]
                })

    for info in shapes_with_tf:
        if info['h'] < 0.6 and info['line_count'] >= 2 and info['max_size'] >= 14:
            old_h = info['h']
            needed_h = info['line_count'] * (info['max_size'] / 72 + 0.05) + 0.1
            new_h = max(old_h, needed_h)
            if new_h > old_h:
                print(f"  第{page_num}页: 摘要文本框 h={old_h:.2f}→{new_h:.2f} (lines={info['line_count']}, size={info['max_size']}pt) '{info['text'][:30]}...'")
                info['shape'].height = Inches(new_h)
                old_bottom = info['top'] + old_h
                new_bottom = info['top'] + new_h
                for other in shapes_with_tf:
                    if other is info:
                        continue
                    if abs(other['left'] - info['left']) < 1.0 and abs(other['top'] - old_bottom) < 0.15:
                        shift = new_bottom - old_bottom
                        new_top = other['top'] + shift
                        print(f"    调整下方文本框 top={other['top']:.2f}→{new_top:.2f} '{other['text'][:30]}...'")
                        other['shape'].top = Inches(new_top)

for idx, slide in enumerate(prs.slides):
    page = idx + 1
    if page in [11, 12, 13, 15]:
        print(f"\n=== 修复第 {page} 页 ===")
        fix_overlap(slide, page)

# === 第三部分：新增 v4.0.6 + v4.0.7 变更页 ===
COLOR_TITLE = RGBColor(0x21, 0x21, 0x21)
COLOR_ACCENT = RGBColor(0x21, 0x96, 0xF3)
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)

def add_changelog_slide(prs, title, changes):
    blank_layout = prs.slide_layouts[6]
    new_slide = prs.slides.add_slide(blank_layout)

    txBox = new_slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE
    run.font.name = "Microsoft YaHei"

    content_box = new_slide.shapes.add_textbox(
        Inches(0.5), Inches(1.15), Inches(12.3), Inches(6.1)
    )
    ctf = content_box.text_frame
    ctf.word_wrap = True
    ctf.auto_size = MSO_AUTO_SIZE.NONE

    for idx, (text, is_bold, is_gap) in enumerate(changes):
        p = ctf.paragraphs[0] if idx == 0 else ctf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.0
        if is_gap:
            p.space_before = Pt(4)
            p.space_after = Pt(0)
            run = p.add_run()
            run.text = ""
            run.font.size = Pt(6)
        else:
            p.space_before = Pt(0)
            p.space_after = Pt(2)
            run = p.add_run()
            run.text = text
            run.font.size = Pt(14) if is_bold else Pt(11)
            run.font.bold = is_bold
            run.font.color.rgb = COLOR_ACCENT if is_bold else COLOR_TEXT
            run.font.name = "Microsoft YaHei"

v406_changes = [
    ("1. PPT 第11/12/13/15页文字重叠修复", True, False),
    ("   增大摘要文本框高度 + 调整下方文本框位置，消除文字重叠", False, False),
    ("", False, True),
    ("2. 命名规则增加「序号」字段 + 水印增加序号", True, False),
    ("   NameSegment 新增 SERIAL 枚举，下拉菜单显示 5 个选项", False, False),
    ("   水印在 serial 非空时显示序号段（4段：日期/序号/地址/经纬度）", False, False),
    ("", False, True),
    ("3. 广角显示改为「广角:开」/「广角:关」", True, False),
    ("   按钮文案直观化，激活/未激活状态明确", False, False),
    ("", False, True),
    ("4. 广角开关与缩放拉杆平行并排", True, False),
    ("   同一 Row 内水平排列，按钮 36dp + 间距 8dp，确保华为 mate70 不溢出", False, False),
    ("", False, True),
    ("5. 闪光灯功能（关闭/自动/常亮/开启）", True, False),
    ("   拍摄页新增 4 选项闪光灯控制，常亮=torch 持续亮，开启=拍照瞬间闪", False, False),
    ("", False, True),
    ("6. 添加查勘条目按钮上移", True, False),
    ("   与搜索框间距缩短至 4dp，列表获得更多垂直空间", False, False),
    ("", False, True),
    ("7. 图片导出功能", True, False),
    ("   按序号/地址（模糊匹配）/借款人三维度筛选", False, False),
    ("   导出 PDF（合并为一个文件，每张一页）或 zip 压缩包", False, False),
    ("", False, True),
    ("8. 双版本生成 + A版离线设备绑定", True, False),
    ("   A版(trial)：Android ID 绑定 + 有效期至 2026-12-31，到期完全锁定", False, False),
    ("   B版(full)：无限制，直接使用", False, False),
    ("   设备识别码采用 Android ID（无需权限、全版本通用）", False, False),
    ("   设置页「关于」显示设备识别码，方便用户查看并告知作者激活", False, False),
]

v407_changes = [
    ("1. 切换为 Release 构建 + R8 full mode 混淆", True, False),
    ("   分发版本从 debug 改为 release，代码混淆 + 资源压缩，提升逆向难度", False, False),
    ("", False, True),
    ("2. 运行时完整性校验（签名校验 + 反调试 + Frida 检测）", True, False),
    ("   签名证书 SHA-256 比对，防止重打包重新签名", False, False),
    ("   调试器检测（Debug.isDebuggerConnected + TracerPid）+ Frida 端口 27042 探测", False, False),
    ("   Application.onCreate + MainActivity.onCreate 双点校验，失败显示锁定页", False, False),
    ("", False, True),
    ("3. 授权设备 ID 改为 SHA-256 哈希存储", True, False),
    ("   BuildConfig 注入设备 ID 的 SHA-256 哈希，原始设备 ID 不再出现在 dex 中", False, False),
    ("   LicenseChecker 运行时对 Android ID 做哈希后比对", False, False),
    ("", False, True),
    ("4. 导出文件支持保存到本地", True, False),
    ("   导出完成后显示结果弹窗：分享 / 保存到本地 / 关闭", False, False),
    ("   分享通过 ACTION_SEND 选择器（微信/邮件/钉钉等已安装应用）", False, False),
    ("   保存到本地通过 SAF（ACTION_CREATE_DOCUMENT）让用户选择目标位置", False, False),
    ("", False, True),
    ("5. 恢复出厂设置警告说明", True, False),
    ("   LockScreen 未授权状态显示「重要说明」警告卡片", False, False),
    ("   SettingsScreen 关于卡片（trial 版）显示同样警告", False, False),
    ("   说明：恢复出厂将改变设备识别码导致 App 不可用，有效期内仅一次重新激活机会", False, False),
    ("", False, True),
    ("6. 构建方式更新", True, False),
    ("   新增 release keystore 签名配置，版本号 versionCode=8 / versionName=4.0.7", False, False),
    ("   构建命令：assembleFullRelease（B版）+ assembleTrialRelease（A版）", False, False),
]

add_changelog_slide(prs, "v4.0.6 更新内容（2026-07-08）", v406_changes)
add_changelog_slide(prs, "v4.0.7 更新内容（2026-07-08）", v407_changes)

prs.save(DST)
print(f"\nPPT saved: {DST}")
print(f"Total slides: {len(prs.slides)}")
